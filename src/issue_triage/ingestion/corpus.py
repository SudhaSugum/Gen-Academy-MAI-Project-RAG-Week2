"""Load the synthetic source of truth from its intended office document formats."""

from __future__ import annotations

import csv
from functools import lru_cache
from pathlib import Path
from typing import Any
from datetime import date

from docx import Document as WordDocument
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[3]
DATA = ROOT / "data" / "synthetic"


def _csv_rows(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def _base_metadata(record: dict[str, str], source_type: str, doc_id: str, chunk_id: str, title: str) -> dict[str, Any]:
    keys = ("client_id", "audit_id", "issue_id", "status", "final_severity", "root_cause", "control_id", "control_version", "outcome", "issue_date", "closure_date", "rau_id", "attribute", "effective_date", "section", "source")
    metadata: dict[str, Any] = {key: str(record[key]) for key in keys if record.get(key) not in (None, "")}
    if record.get("issue_date"):
        # Pinecone range operators require numeric metadata; keep the ISO date
        # for display and add a sortable YYYYMMDD key for lookback filters.
        metadata["issue_date_ordinal"] = int(date.fromisoformat(record["issue_date"]).strftime("%Y%m%d"))
    metadata.update({"source_type": source_type, "doc_id": doc_id, "chunk_id": chunk_id, "title": title})
    return metadata


@lru_cache(maxsize=1)
def build_documents() -> dict[str, list[Document]]:
    documents: dict[str, list[Document]] = {"past_issues": [], "controls": [], "policies": []}
    for row in _csv_rows(DATA / "past_issues.csv"):
        doc_id = row["doc_id"]
        date_lines = []
        if row.get("issue_date"):
            date_lines.append(f"Issue date: {row['issue_date']}")
        if row.get("closure_date"):
            date_lines.append(f"Closure date: {row['closure_date']}")
        date_text = "\n".join(date_lines)
        body = f"Issue: {row['title']}\nAudit ID: {row['audit_id']}\nDescription: {row['description']}\n{date_text}\nFinal severity: {row['final_severity']}\nRoot cause: {row['root_cause']}\nControl ID: {row['control_id']}\nRemediation: {row['remediation']}\nRecorded outcome: {row['outcome']}\nStatus: {row['status']}"
        metadata = _base_metadata(row, "past_issue", doc_id, f"{doc_id}-chunk-001", row["title"])
        documents["past_issues"].append(Document(page_content=body, metadata=metadata))

    for path in sorted((DATA / "controls_docx").glob("*.docx")):
        word = WordDocument(path)
        text = "\n".join(paragraph.text.strip() for paragraph in word.paragraphs if paragraph.text.strip())
        # Tables are laid out as label/value cells; preserve the source document's text as citation evidence.
        values: dict[str, str] = {}
        for table in word.tables:
            for row in table.rows:
                if len(row.cells) >= 2:
                    values[row.cells[0].text.strip()] = row.cells[1].text.strip()
        doc_id = values.get("Document ID", path.stem.upper())
        control_id = values.get("Control ID", "")
        title = next((p.text.strip() for p in word.paragraphs if p.style.name == "Title"), control_id)
        text = "\n".join([text, *(f"{key}: {value}" for key, value in values.items()), *[f"{p.text.strip()}" for p in word.paragraphs if p.text.strip() and p.style.name != "Title"]])
        record = {"client_id": values.get("Client", ""), "audit_id": values.get("Audit", ""), "control_id": control_id, "control_version": values.get("Version", ""), "rau_id": f"RAU-{control_id[:2]}-01", "attribute": values.get("Risk and control attribute", ""), "effective_date": values.get("Effective date", ""), "source": "Synthetic controls library"}
        metadata = _base_metadata(record, "control", doc_id, f"{doc_id}-chunk-001", title)
        documents["controls"].append(Document(page_content=text, metadata=metadata))

    splitter = RecursiveCharacterTextSplitter(chunk_size=650, chunk_overlap=60, add_start_index=True)
    pdf_dir = DATA / "policies_pdf"
    policy_specs = {path.stem.upper(): path for path in (DATA / "policies").glob("*.md")}
    for path in sorted(pdf_dir.glob("*.pdf")):
        policy_id = path.stem.upper()
        reader = PdfReader(str(path))
        content = "\n".join(page.extract_text() or "" for page in reader.pages).strip()
        md_path = policy_specs.get(policy_id)
        headings = [line.removeprefix("# ").strip() for line in md_path.read_text(encoding="utf-8").splitlines() if line.startswith("# ")] if md_path else []
        section = headings[0] if headings else policy_id
        for client_id, audit_id in (("synthetic-client-a", "synthetic-audit-iam-2026"), ("synthetic-client-b", "synthetic-audit-finance-2026")):
            tenant_text = f"Client: {client_id}\nAudit: {audit_id}\n{content}"
            for i, chunk in enumerate(splitter.split_text(tenant_text), 1):
                doc_id = f"SYN-POL-{policy_id}-{client_id[-1]}"
                record = {"client_id": client_id, "audit_id": audit_id, "section": section, "effective_date": "2026-01-01", "source": "Synthetic policy library"}
                metadata = _base_metadata(record, "policy", doc_id, f"{doc_id}-chunk-{i:03d}", section)
                documents["policies"].append(Document(page_content=chunk, metadata=metadata))
    return documents
