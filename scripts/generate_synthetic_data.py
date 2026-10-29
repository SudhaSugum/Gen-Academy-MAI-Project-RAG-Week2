"""Generate the deterministic, fictional V1 audit corpus (no client data)."""

from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

ROOT = Path(__file__).resolve().parents[1] / "data" / "synthetic"
ROOT.mkdir(parents=True, exist_ok=True)

CONTROL_SPECS = [
    ("AC-01", "Unique user identity", "Every person is assigned a distinct account linked to an approved identity record.", "authorization"),
    ("AC-02", "Access authorization", "A designated manager authorizes access before it is provisioned, and the approval trail is retained.", "authorization"),
    ("AC-03", "Role-based access", "Users receive access through approved roles, with conflicting responsibilities separated.", "segregation of duties"),
    ("AC-04", "Periodic access recertification", "At least quarterly, application owners recertify assigned accounts, confirm continued business need, and retain dated sign-off.", "completeness"),
    ("AC-05", "Privileged access approval", "Elevated access is granted only with recorded owner authorization and a stated business justification.", "authorization"),
    ("AC-06", "Privileged activity monitoring", "Privileged operations are logged and periodically inspected by someone independent of the activity.", "monitoring"),
    ("AC-07", "Timely deprovisioning", "Accounts are disabled soon after employment ends or another event removes the need for access.", "timeliness"),
    ("AC-08", "Joiner mover leaver reconciliation", "HR joiner, mover, and leaver updates are reconciled to identity-account changes, with exceptions tracked to closure.", "completeness"),
    ("AC-09", "Emergency access review", "Emergency access expires promptly, is monitored during use, and receives an independent post-event review.", "monitoring"),
    ("CM-01", "Change authorization", "A production release must have traceable approval before it is deployed.", "authorization"),
    ("CM-02", "Change testing", "Teams verify each change against documented acceptance tests before release.", "accuracy"),
    ("CM-03", "Change segregation", "Development, approval, and production deployment are separated so a developer cannot authorize and release their own change.", "segregation of duties"),
    ("PR-01", "Supplier onboarding", "A supplier must pass a risk review and due diligence and receive approval before onboarding.", "authorization"),
    ("PR-02", "Supplier access review", "Third-party accounts are periodically recertified and disabled when the supplier's business need ends.", "completeness"),
    ("FN-01", "Journal entry approval", "Manual journal entries receive an independent review and authorization before posting.", "authorization"),
    ("FN-02", "Account reconciliation", "Reconciliations are performed each month for key accounts; aged differences are assigned and escalated.", "completeness"),
    ("FN-03", "Payment release", "Two authorized approvers release payment batches according to approved value limits.", "segregation of duties"),
    ("FN-04", "Master data changes", "Sensitive reference-data updates need documented authorization and an independent accuracy check.", "accuracy"),
]

ISSUE_PATTERNS = [
    {"control":"AC-04", "title":"Quarterly application access attestation missed its deadline", "description":"An application owner did not finish the scheduled account recertification by quarter close. Sign-off was missing for several sampled accounts, leaving continued business need unconfirmed.", "severity":"Medium", "root_cause":"Monitoring and escalation", "remediation":"Assign the review to a named process owner, complete a look-back of affected accounts, and add a deadline alert to the quarterly attestation tracker.", "outcome":"The owner reviewed the outstanding accounts; the next quarterly cycle completed on schedule with dated sign-offs.", "status":"closed", "attribute":"completeness"},
    {"control":"AC-07", "title":"Former staff account disablement was not traceable", "description":"For a few departed workers, the audit could not establish when system access was shut off because the HR-to-identity handoff record was missing.", "severity":"High", "root_cause":"Access lifecycle handoff", "remediation":"Confirm the current status of sampled accounts, retain identity-system logs, and reconcile departure notices to disablement timestamps.", "outcome":"Most sampled accounts were shown disabled in system logs; an account without confirmation was escalated and shut down.", "status":"closed", "attribute":"timeliness"},
    {"control":"AC-02", "title":"Provisioning preceded manager authorization", "description":"Some sampled access requests showed that a manager's sign-off occurred after the account change. The review found no evidence that the access was unauthorized.", "severity":"Medium", "root_cause":"Process execution", "remediation":"Reinforce the approval-before-access step and produce a timestamp exception listing for the access owner to review.", "outcome":"The approval sequence was corrected, and a later sample showed authorization recorded before provisioning.", "status":"closed", "attribute":"authorization"},
    {"control":"CM-01", "title":"Production deployment had no recorded pre-approval", "description":"One sampled production release lacked an approval record from before implementation. A follow-up review found no identified adverse effect.", "severity":"Medium", "root_cause":"Process execution", "remediation":"Make pre-release approval a required workflow entry and have the change owner review emergency exceptions each month.", "outcome":"Workflow checks were added, and the next monthly sample contained complete approvals.", "status":"closed", "attribute":"authorization"},
    {"control":"CM-03", "title":"One operator developed and released a production change", "description":"A sampled release listed the same operator as developer and production deployer. The file contained no evidence of a separate post-release review.", "severity":"High", "root_cause":"Technology configuration", "remediation":"Separate release permissions, retrospectively inspect affected deployments, and alert when conflicting operator accounts are detected.", "outcome":"Conflicting permissions were removed and release monitoring was put in place.", "status":"closed", "attribute":"segregation of duties"},
    {"control":"FN-02", "title":"Aged reconciliation items lacked follow-up", "description":"The monthly close included old reconciling items without a recorded escalation path or target resolution date.", "severity":"Medium", "root_cause":"Monitoring and escalation", "remediation":"Define aging triggers, assign accountable roles to open items, and track each exception through documented resolution.", "outcome":"Owners were assigned to the aged items, and the next two close cycles showed timely escalation.", "status":"closed", "attribute":"completeness"},
]

clients = [("synthetic-client-a", "synthetic-audit-iam-2026"), ("synthetic-client-b", "synthetic-audit-finance-2026")]

with (ROOT / "controls.csv").open("w", newline="", encoding="utf-8") as f:
    fields = ["doc_id","title","text","client_id","audit_id","control_id","control_version","rau_id","attribute","effective_date","source"]
    writer = csv.DictWriter(f, fieldnames=fields); writer.writeheader()
    for client, audit in clients:
        for control_id, name, description, attribute in CONTROL_SPECS:
            writer.writerow({"doc_id":f"SYN-CTRL-{client[-1]}-{control_id}","title":f"{control_id} — {name}","text":f"Synthetic example control {control_id}: {description} Control owner role: Process Owner. Exceptions are tracked to resolution.","client_id":client,"audit_id":audit,"control_id":control_id,"control_version":"1.0","rau_id":f"RAU-{control_id[:2]}-01","attribute":attribute,"effective_date":"2026-01-01","source":"Synthetic controls library"})

with (ROOT / "past_issues.csv").open("w", newline="", encoding="utf-8") as f:
    fields = ["doc_id","title","description","client_id","audit_id","issue_id","status","final_severity","root_cause","control_id","control_version","remediation","outcome","issue_date","closure_date","rau_id","attribute","source"]
    writer = csv.DictWriter(f, fieldnames=fields); writer.writeheader()
    for i in range(60):
        pattern = ISSUE_PATTERNS[i % len(ISSUE_PATTERNS)]
        client, audit = clients[(i // 2) % len(clients)]
        # Keep the corpus tenant/audit boundary consistent with the control catalogue.
        if pattern["control"].startswith(("AC", "CM")):
            client, audit = clients[0]
        else:
            client, audit = clients[1]
        issue_num = i + 1
        # Keep two audit periods per client so cross-audit precedent and
        # exact-audit duplicate filtering can be exercised with synthetic data.
        if (i // len(ISSUE_PATTERNS)) % 2 == 1:
            audit = f"{audit.rsplit('-', 1)[0]}-2025"
        # Keep most examples inside the default two-year scoping window while
        # retaining a few older records for the excluded-issue display.
        issue_year = "2024" if issue_num % 4 == 0 else "2025"
        title_variants = [pattern["title"], f"Evidence gap: {pattern['title'].lower()}", f"Review noted {pattern['title'].lower()}"]
        writer.writerow({
            "doc_id":f"SYN-ISS-{issue_num:03d}","title":title_variants[i % len(title_variants)],
            "description":f"Synthetic scenario {issue_num:03d}. {pattern['description']} Scope was limited to the fictional {client} engagement sample.",
            "client_id":client,"audit_id":audit,"issue_id":f"SYN-ISS-{issue_num:03d}","status":("open" if issue_num % 11 == 0 else pattern["status"]),
            "final_severity":pattern["severity"],"root_cause":pattern["root_cause"],"control_id":pattern["control"],"control_version":"1.0",
            "remediation":pattern["remediation"],"outcome":pattern["outcome"],"issue_date":f"{issue_year}-{(issue_num%12)+1:02d}-15",
            "closure_date":"" if issue_num % 11 == 0 else f"{int(issue_year)+1}-{(issue_num%12)+1:02d}-20",
            "rau_id":f"RAU-{pattern['control'][:2]}-01","attribute":pattern["attribute"],"source":"Synthetic audit issue register",
        })

policy_specs = [
    ("SEC-01", "Identity lifecycle", "Synthetic policy example. Grant an account only after approval, update its access when responsibilities shift, and disable it promptly when the need ends. Reconcile HR joiner, mover, and leaver notices with identity-account changes."),
    ("SEC-02", "Privileged access", "Synthetic policy example. Elevated rights need a recorded business justification and bounded authorization. Log privileged activity and arrange an independent review on a recurring schedule."),
    ("FIN-01", "Financial close", "Synthetic policy example. Complete key balance reconciliations each month and obtain an independent review. Assign owners to aged items and escalate them for resolution."),
    ("CHG-01", "Production change management", "Synthetic policy example. Record authorization before a production release, verify it against documented acceptance tests, and keep development, approval, and deployment duties separate."),
    ("TPR-01", "Third-party risk", "Synthetic policy example. Complete supplier due diligence and approval before onboarding. Reassess external-party access periodically, monitor its use, and disable it when the engagement no longer needs it."),
]
policy_dir = ROOT / "policies"; policy_dir.mkdir(exist_ok=True)
for policy_id, title, body in policy_specs:
    (policy_dir / f"{policy_id.lower()}.md").write_text(f"# {policy_id} — {title}\n\n{body}\n", encoding="utf-8")

controls_dir = ROOT / "controls_docx"
controls_dir.mkdir(exist_ok=True)
for control_id, name, description, attribute in CONTROL_SPECS:
    for client, audit in clients:
        doc_id = f"SYN-CTRL-{client[-1]}-{control_id}"
        doc = Document()
        doc.add_heading(f"{control_id} - {name}", 0)
        doc.add_paragraph("SYNTHETIC CONTROL RECORD", style="Subtitle")
        table = doc.add_table(rows=0, cols=2)
        table.style = "Light Shading Accent 1"
        for label, value in (("Control ID", control_id), ("Document ID", doc_id), ("Client", client), ("Audit", audit), ("Version", "1.0"), ("Effective date", "2026-01-01"), ("Risk and control attribute", attribute)):
            cells = table.add_row().cells
            cells[0].text, cells[1].text = label, value
        doc.add_heading("Control description", level=1)
        doc.add_paragraph(f"Synthetic example control {control_id}: {description}")
        doc.add_heading("Control owner role", level=1)
        doc.add_paragraph("Process Owner")
        doc.add_heading("Exception handling", level=1)
        doc.add_paragraph("Exceptions are tracked to resolution.")
        doc.add_paragraph("Source: Synthetic controls library")
        doc.core_properties.title = f"{control_id} - {name}"
        doc.core_properties.subject = "Synthetic audit control"
        doc.save(controls_dir / f"{doc_id.lower()}.docx")

policies_pdf_dir = ROOT / "policies_pdf"
policies_pdf_dir.mkdir(exist_ok=True)
styles = getSampleStyleSheet()
styles.add(ParagraphStyle(name="PolicyTitle", parent=styles["Title"], textColor=colors.HexColor("#15324B"), spaceAfter=16))
styles.add(ParagraphStyle(name="PolicyBody", parent=styles["BodyText"], leading=16, fontSize=11, spaceAfter=10))
styles.add(ParagraphStyle(name="PolicyEyebrow", parent=styles["Normal"], textColor=colors.HexColor("#52728A"), fontName="Helvetica-Bold", fontSize=9, spaceAfter=8))
for policy_id, title, body in policy_specs:
    output = policies_pdf_dir / f"{policy_id.lower()}.pdf"
    story = [Paragraph("SYNTHETIC POLICY", styles["PolicyEyebrow"]), Paragraph(f"{policy_id} - {title}", styles["PolicyTitle"])]
    meta = Table([["Policy ID", policy_id], ["Effective date", "2026-01-01"], ["Source", "Synthetic policy library"]], colWidths=[1.4 * inch, 4.9 * inch])
    meta.setStyle(TableStyle([("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EAF0F5")), ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#15324B")), ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CAD4DE")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("PADDING", (0, 0), (-1, -1), 8)]))
    story.extend([meta, Spacer(1, 24), Paragraph("Policy statement", styles["Heading2"]), Paragraph(body, styles["PolicyBody"]), Spacer(1, 12), Paragraph("This fictional policy is for local concept-build use only.", styles["Italic"])])
    SimpleDocTemplate(str(output), pagesize=letter, rightMargin=0.8*inch, leftMargin=0.8*inch, topMargin=0.7*inch, bottomMargin=0.7*inch, title=f"{policy_id} - {title}", author="Synthetic policy library").build(story)

print(f"Generated 60 synthetic issues, 36 synthetic control Word files, and {len(policy_specs)} synthetic policy PDFs in {ROOT}")
